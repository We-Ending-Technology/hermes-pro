from __future__ import annotations

import io
import json
import logging
from urllib.parse import quote

import httpx
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from backend.app.ai_gateway.factory import build_ai_gateway
from backend.app.commercial import build_commercial_metadata, mark_publication_ready
from backend.app.core.config import get_settings
from backend.app.db.supabase import SupabaseREST
from backend.app.queue import JobQueue
from backend.app.services.persistence import PersistentStore

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("hermes.worker")

MAX_ATTEMPTS = 3


def _parse_json(text: str) -> dict:
    cleaned = text.strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.strip("`")
        if cleaned.startswith("json"):
            cleaned = cleaned[4:]
    try:
        value = json.loads(cleaned)
    except json.JSONDecodeError as exc:
        raise RuntimeError("A IA retornou JSON inválido para o pipeline") from exc
    if not isinstance(value, dict):
        raise RuntimeError("A IA retornou uma estrutura inválida para o pipeline")
    return value


def _ebook_text(title: str, writer_data: dict) -> list[tuple[str, str]]:
    sections: list[tuple[str, str]] = [(title, writer_data.get("introduction", ""))]
    for chapter in writer_data.get("chapters", []):
        if isinstance(chapter, dict):
            sections.append((str(chapter.get("title", "Capítulo")), str(chapter.get("content", ""))))
    sections.append(("Conclusão", str(writer_data.get("conclusion", ""))))
    return [(heading, body) for heading, body in sections if body.strip()]


def _make_pdf(title: str, writer_data: dict) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=2 * cm, leftMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    story = []
    for heading, body in _ebook_text(title, writer_data):
        story.append(Paragraph(heading, styles["Title"] if heading == title else styles["Heading1"]))
        story.append(Spacer(1, 0.35 * cm))
        for paragraph in body.split("\n"):
            if paragraph.strip():
                story.append(Paragraph(paragraph.strip(), styles["BodyText"]))
                story.append(Spacer(1, 0.18 * cm))
    doc.build(story)
    return buffer.getvalue()


def _make_docx(title: str, writer_data: dict) -> bytes:
    document = Document()
    document.add_heading(title, level=0)
    for heading, body in _ebook_text(title, writer_data)[1:]:
        document.add_heading(heading, level=1)
        for paragraph in body.split("\n"):
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def _make_cover(title: str, topic: str) -> bytes:
    prompt = quote(f"professional ebook cover, modern editorial design, title: {title}, topic: {topic}, no logos, no watermark, clean typography, portrait 2:3")
    url = f"https://image.pollinations.ai/prompt/{prompt}?width=1200&height=1800&nologo=true"
    async with httpx.AsyncClient(timeout=120, follow_redirects=True) as client:
        response = await client.get(url)
        response.raise_for_status()
        if not response.content:
            raise RuntimeError("O provedor de capa retornou um arquivo vazio")
        return response.content


async def process_product(job_id: str, store: PersistentStore, queue: JobQueue) -> None:
    job = await store.get_job(job_id)
    if not job or job["job_type"] != "product_generation" or job["status"] == "completed":
        return
    attempts = int(job["attempts"]) + 1
    product_id = (job.get("payload") or {}).get("product_id")
    topic = (job.get("payload") or {}).get("topic")
    if not product_id or not topic:
        await store.update_job(job_id, "failed", attempts=attempts, error_message="invalid product_generation payload")
        return
    if attempts > MAX_ATTEMPTS:
        await store.update_job(job_id, "failed", attempts=attempts, error_message="maximum attempts reached")
        return

    try:
        await store.update_job(job_id, "running", attempts=attempts, error_message=None)
        settings = get_settings()
        gateway = build_ai_gateway(settings)
        await store.update_product(product_id, status="generating", current_stage="topic")

        strategist = await gateway.complete(
            f"Create a concise product strategy for the digital ebook topic: {topic}. Return valid JSON with keys: positioning, audience, title, subtitle, category, outline (array of chapter titles).",
            system="You are the Hermes Pro strategist. Return only valid JSON. Do not claim market data you do not have.",
        )
        strategy_data = _parse_json(strategist.content)
        title = str(strategy_data.get("title") or topic).strip()
        await store.update_product(product_id, current_stage="writer", title=title, metadata={"strategy": strategy_data, "ai_provider": strategist.provider, "ai_model": strategist.model})

        writer = await gateway.complete(
            f"Using this strategy, write a complete practical ebook for topic '{topic}'. Strategy: {json.dumps(strategy_data, ensure_ascii=False)}. Return valid JSON with keys: introduction, chapters (array of objects with title and content), conclusion. Write original useful content; do not fabricate citations or statistics. Make every chapter substantial and actionable.",
            system="You are the Hermes Pro writer. Return only valid JSON. Make the ebook useful enough to sell honestly, with clear sections, practical examples, checklists where useful, and no invented facts.",
        )
        writer_data = _parse_json(writer.content)
        await store.update_product(product_id, current_stage="reviewer")

        reviewer = await gateway.complete(
            f"Review this ebook for topic '{topic}'. Return valid JSON with keys score (0-100), findings (array), approved (boolean). Approval requires score >= 80. EBOOK: {json.dumps(writer_data, ensure_ascii=False)}",
            system="You are the Hermes Pro reviewer. Be conservative and identify factual, structural and readability problems. Do not invent external facts.",
        )
        review_data = _parse_json(reviewer.content)
        score = int(review_data.get("score", 0))
        approved = bool(review_data.get("approved")) and score >= 80
        metadata = {"strategy": strategy_data, "content": writer_data, "review": review_data, "quality_score": score, "ai_provider": strategist.provider, "ai_model": strategist.model}
        await store.update_product(product_id, current_stage="validation", metadata=metadata)
        if not approved:
            await store.update_product(product_id, status="review_required", current_stage="reviewer", metadata=metadata)
            await store.update_job(job_id, "completed", attempts=attempts)
            return

        await store.update_product(product_id, current_stage="commercial")
        commercial = await gateway.complete(
            f"Create commercial metadata for this ebook. Topic: '{topic}'. Strategy: {json.dumps(strategy_data, ensure_ascii=False)}. Return valid JSON with keys: description, short_description, keywords (array of 5-10 strings), suggested_price (number in BRL), price_rationale. The price is a recommendation only. If you do not have verified current market comparables, say so explicitly in price_rationale and do not invent competitors, prices, sales or market statistics.",
            system="You are the Hermes Pro commercial strategist. Return only valid JSON. Create persuasive but truthful copy. Never invent market evidence. Suggested price must be positive and reasonable for a Brazilian digital ebook.",
        )
        commercial_data = _parse_json(commercial.content)
        commercial_metadata = build_commercial_metadata(
            strategy=strategy_data,
            description=str(commercial_data.get("description") or "").strip(),
            short_description=str(commercial_data.get("short_description") or "").strip(),
            keywords=list(commercial_data.get("keywords") or []),
            suggested_price=float(commercial_data.get("suggested_price") or 0),
            price_rationale=str(commercial_data.get("price_rationale") or "").strip(),
        )
        metadata.update(commercial_metadata)
        await store.update_product(product_id, title=title, metadata=metadata)

        pdf_bytes = _make_pdf(title, writer_data)
        docx_bytes = _make_docx(title, writer_data)
        await store.update_product(product_id, current_stage="document")
        pdf_path = f"{product_id}/ebook.pdf"
        docx_path = f"{product_id}/ebook.docx"
        pdf_url = await store.db.upload("ebooks", pdf_path, pdf_bytes, "application/pdf")
        docx_url = await store.db.upload("exports", docx_path, docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

        await store.update_product(product_id, current_stage="cover")
        cover_bytes = await _make_cover(title, topic)
        cover_path = f"{product_id}/cover.jpg"
        cover_url = await store.db.upload("covers", cover_path, cover_bytes, "image/jpeg")

        metadata.update({
            "pdf_url": pdf_url,
            "document_url": docx_url,
            "cover_url": cover_url,
            "document_path": f"ebooks/{pdf_path}",
            "docx_path": f"exports/{docx_path}",
            "cover_path": f"covers/{cover_path}",
            "asset_status": "complete",
        })
        metadata = mark_publication_ready(metadata)
        await store.update_product(product_id, status="ready_to_sell" if metadata["publication_ready"] else "completed", current_stage="publication", title=title, metadata=metadata)
        await store.update_job(job_id, "completed", attempts=attempts)
        logger.info("completed full commercial product job=%s product=%s", job_id, product_id)
    except Exception as exc:
        message = str(exc)[:1000]
        if attempts < MAX_ATTEMPTS:
            await store.update_job(job_id, "retrying", attempts=attempts, error_message=message)
            await queue.enqueue(job_id)
        else:
            await store.update_job(job_id, "failed", attempts=attempts, error_message=message)
            await store.update_product(product_id, status="failed")
        logger.exception("product_generation failed job=%s", job_id)


async def recover_pending(store: PersistentStore, queue: JobQueue) -> None:
    for job in await store.list_jobs():
        if job["job_type"] == "product_generation" and job["status"] in {"pending", "retrying"}:
            await queue.enqueue(str(job["id"]))


async def run() -> None:
    settings = get_settings()
    store = PersistentStore(SupabaseREST(settings))
    queue = JobQueue(settings.redis_url)
    logger.info("Hermes Pro worker started; queue=hermes:jobs:product_generation")
    try:
        await recover_pending(store, queue)
        while True:
            job_id = await queue.dequeue(timeout=10)
            if job_id:
                await process_product(job_id, store, queue)
            else:
                await recover_pending(store, queue)
    finally:
        await queue.close()


if __name__ == "__main__":
    import asyncio
    try:
        asyncio.run(run())
    except KeyboardInterrupt:
        logger.info("Hermes Pro worker stopped")
