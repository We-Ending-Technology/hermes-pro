from __future__ import annotations

import io
import json
import re
from typing import Any

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from PIL import Image, ImageDraw, ImageFont


def normalize_content(raw: str | dict[str, Any]) -> dict[str, Any]:
    if isinstance(raw, dict):
        data = raw
    else:
        try:
            data = json.loads(raw)
        except (TypeError, ValueError):
            data = {"introduction": raw, "chapters": [], "conclusion": ""}
    chapters = data.get("chapters") or []
    normalized = []
    for index, chapter in enumerate(chapters, 1):
        if isinstance(chapter, dict):
            normalized.append({"title": str(chapter.get("title") or f"Capítulo {index}"), "content": str(chapter.get("content") or "")})
        else:
            normalized.append({"title": f"Capítulo {index}", "content": str(chapter)})
    return {
        "introduction": str(data.get("introduction") or ""),
        "chapters": normalized,
        "conclusion": str(data.get("conclusion") or ""),
    }


def build_docx(title: str, content: dict[str, Any]) -> bytes:
    doc = Document()
    doc.add_heading(title, 0)
    if content["introduction"]:
        doc.add_heading("Introdução", level=1)
        doc.add_paragraph(content["introduction"])
    for chapter in content["chapters"]:
        doc.add_heading(chapter["title"], level=1)
        for paragraph in re.split(r"\n\s*\n", chapter["content"]):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())
    if content["conclusion"]:
        doc.add_heading("Conclusão", level=1)
        doc.add_paragraph(content["conclusion"])
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def build_pdf(title: str, content: dict[str, Any]) -> bytes:
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=20 * mm, leftMargin=20 * mm, topMargin=20 * mm, bottomMargin=20 * mm)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 8 * mm)]
    if content["introduction"]:
        story += [Paragraph("Introdução", styles["Heading1"]), Paragraph(content["introduction"].replace("\n", "<br/>") , styles["BodyText"]), Spacer(1, 4 * mm)]
    for chapter in content["chapters"]:
        story += [Paragraph(chapter["title"], styles["Heading1"]), Paragraph(chapter["content"].replace("\n", "<br/>") , styles["BodyText"]), Spacer(1, 4 * mm)]
    if content["conclusion"]:
        story += [Paragraph("Conclusão", styles["Heading1"]), Paragraph(content["conclusion"].replace("\n", "<br/>") , styles["BodyText"])]
    doc.build(story)
    return output.getvalue()


def build_cover(title: str, subtitle: str = "") -> bytes:
    width, height = 1600, 2560
    image = Image.new("RGB", (width, height), (18, 18, 18))
    draw = ImageDraw.Draw(image)
    try:
        font_title = ImageFont.truetype("DejaVuSans-Bold.ttf", 110)
        font_subtitle = ImageFont.truetype("DejaVuSans.ttf", 52)
    except OSError:
        font_title = ImageFont.load_default()
        font_subtitle = ImageFont.load_default()
    draw.rectangle((90, 90, width - 90, height - 90), outline=(190, 150, 70), width=6)
    words = title.split()
    lines, current = [], ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=font_title)[2] > width - 260 and current:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    y = height // 2 - len(lines) * 65
    for line in lines:
        box = draw.textbbox((0, 0), line, font=font_title)
        draw.text(((width - (box[2] - box[0])) / 2, y), line, fill=(238, 224, 190), font=font_title)
        y += 135
    if subtitle:
        draw.text((140, height - 360), subtitle[:90], fill=(210, 210, 210), font=font_subtitle)
    output = io.BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue()
