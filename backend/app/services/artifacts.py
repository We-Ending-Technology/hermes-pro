from __future__ import annotations

import html
import io
import re
from pathlib import Path
from typing import Any

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def _safe_name(value: str) -> str:
    value = re.sub(r"[^a-zA-Z0-9_-]+", "-", value.strip().lower()).strip("-")
    return value[:80] or "ebook"


def _font(size: int):
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def _content_data(content: str) -> dict[str, Any]:
    import json
    try:
        data = json.loads(content)
        if isinstance(data, dict):
            return data
    except Exception:
        pass
    return {"introduction": content, "chapters": [], "conclusion": ""}


def build_cover(title: str, subtitle: str) -> bytes:
    width, height = 1600, 2560
    image = Image.new("RGB", (width, height), "#10111b")
    draw = ImageDraw.Draw(image)
    accent = "#8b5cf6"
    draw.rounded_rectangle((90, 90, width - 90, height - 90), radius=50, outline=accent, width=5)
    draw.text((140, 180), "HERMES PRO", fill="#d8d8e8", font=_font(56))
    draw.text((140, 330), "DIGITAL EDITION", fill=accent, font=_font(42))

    def wrapped(text: str, max_chars: int = 24):
        words, lines, current = text.split(), [], ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if len(candidate) > max_chars and current:
                lines.append(current)
                current = word
            else:
                current = candidate
        if current:
            lines.append(current)
        return lines

    y = 760
    for line in wrapped(title, 24):
        draw.text((140, y), line, fill="#ffffff", font=_font(104))
        y += 125
    y += 80
    for line in wrapped(subtitle, 38):
        draw.text((140, y), line, fill="#b7b8c8", font=_font(48))
        y += 68
    draw.ellipse((width - 420, height - 520, width - 180, height - 280), fill=accent)
    draw.text((width - 370, height - 470), "H", fill="#ffffff", font=_font(130))

    output = io.BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue()


def build_docx(title: str, subtitle: str, content: str) -> bytes:
    data = _content_data(content)
    doc = Document()
    doc.add_heading(title, 0)
    if subtitle:
        doc.add_paragraph(subtitle)
    if data.get("introduction"):
        doc.add_heading("Introdução", level=1)
        doc.add_paragraph(str(data["introduction"]))
    for chapter in data.get("chapters", []):
        if isinstance(chapter, dict):
            doc.add_heading(str(chapter.get("title", "Capítulo")), level=1)
            doc.add_paragraph(str(chapter.get("content", "")))
    if data.get("conclusion"):
        doc.add_heading("Conclusão", level=1)
        doc.add_paragraph(str(data["conclusion"]))
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def build_pdf(title: str, subtitle: str, content: str) -> bytes:
    data = _content_data(content)
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=48)
    styles = getSampleStyleSheet()
    story = [Paragraph(html.escape(title), styles["Title"])]
    if subtitle:
        story += [Paragraph(html.escape(subtitle), styles["Normal"]), Spacer(1, 18)]
    if data.get("introduction"):
        story += [Paragraph("Introdução", styles["Heading1"]), Paragraph(html.escape(str(data["introduction"])), styles["BodyText"])]
    for chapter in data.get("chapters", []):
        if isinstance(chapter, dict):
            story += [Spacer(1, 12), Paragraph(html.escape(str(chapter.get("title", "Capítulo"))), styles["Heading1"]), Paragraph(html.escape(str(chapter.get("content", ""))), styles["BodyText"])]
    if data.get("conclusion"):
        story += [Spacer(1, 12), Paragraph("Conclusão", styles["Heading1"]), Paragraph(html.escape(str(data["conclusion"])), styles["BodyText"])]
    doc.build(story)
    return output.getvalue()


def build_artifacts(title: str, subtitle: str, content: str) -> dict[str, tuple[str, bytes, str]]:
    slug = _safe_name(title)
    return {
        "cover": (f"{slug}/cover.png", build_cover(title, subtitle), "image/png"),
        "docx": (f"{slug}/ebook.docx", build_docx(title, subtitle, content), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        "pdf": (f"{slug}/ebook.pdf", build_pdf(title, subtitle, content), "application/pdf"),
    }
